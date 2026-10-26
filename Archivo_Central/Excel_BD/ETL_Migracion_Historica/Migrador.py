import os
import glob
import datetime
import tkinter as tk
from tkinter import filedialog
from difflib import get_close_matches
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter, range_boundaries

# ==========================================
# CATÁLOGO DE ÁREAS
# ==========================================
CATALOGO_AREAS = [
    "ADMINISTRACIÓN LOCAL AZUA", "ADMINISTRACIÓN LOCAL BANÍ", "ADMINISTRACIÓN LOCAL BARAHONA",
    "ADMINISTRACIÓN LOCAL BÁVARO", "ADMINISTRACIÓN LOCAL BOCA CHICA", "ADMINISTRACIÓN LOCAL BONAO",
    "ADMINISTRACIÓN LOCAL EL SEIBO", "ADMINISTRACIÓN LOCAL HERRERA", "ADMINISTRACIÓN LOCAL HIGÜEY",
    "ADMINISTRACIÓN LOCAL LA FERIA", "ADMINISTRACIÓN LOCAL LA ROMANA", "ADMINISTRACIÓN LOCAL LA VEGA",
    "ADMINISTRACIÓN LOCAL LINCOLN", "ADMINISTRACIÓN LOCAL LOS MINA", "ADMINISTRACIÓN LOCAL LOS PRÓCERES",
    "ADMINISTRACIÓN LOCAL MÁXIMO GÓMEZ", "ADMINISTRACIÓN LOCAL MOCA", "ADMINISTRACIÓN LOCAL MONTECRISTI",
    "ADMINISTRACIÓN LOCAL NAGUA", "ADMINISTRACIÓN LOCAL OGC SANTIAGO", "ADMINISTRACIÓN LOCAL PUERTO PLATA",
    "ADMINISTRACIÓN LOCAL SALCEDO", "ADMINISTRACIÓN LOCAL SAMANÁ", "ADMINISTRACIÓN LOCAL SAN CARLOS",
    "ADMINISTRACIÓN LOCAL SAN CRISTÓBAL", "ADMINISTRACIÓN LOCAL SAN FRANCISCO DE MACORÍS",
    "ADMINISTRACIÓN LOCAL SAN JUAN DE LA MAGUANA", "ADMINISTRACIÓN LOCAL SAN PEDRO DE MACORÍS",
    "ADMINISTRACIÓN LOCAL SANTIAGO", "ADMINISTRACIÓN LOCAL VALVERDE MAO", "ADMINISTRACIÓN LOCAL VILLA MELLA",
    "ADMINISTRACIÓN LOCAL ZONA ORIENTAL", "CENTRO DE FISCALIZACIÓN LOS PRÓCERES",
    "CENTRO DE FISCALIZACIÓN MÁXIMO GÓMEZ", "ADMINISTRACIÓN LOCAL HATO MAYOR",
    "CENTRO DE SERVICIOS LOCAL COTUÍ", "CENTRO DE SERVICIOS LOCAL LAS TERRENAS",
    "CENTRO DE SERVICIOS MONTE PLATA", "CENTRO DE SERVICIOS VILLA ALTAGRACIA",
    "COLECTURÍA CONSTANZA", "COLECTURÍA DAJABÓN", "COLECTURÍA DE COTUÍ",
    "COLECTURÍA DE SAN JOSÉ DE OCOA", "COLECTURÍA ELÍAS PIÑA", "COLECTURÍA JARABACOA",
    "COLECTURÍA JIMANÍ", "COLECTURÍA LAS MATAS DE FARFÁN", "COLECTURÍA NEYBA",
    "COLECTURÍA PEDERNALES", "COLECTURÍA SAN JOSÉ DE OCOA", "COLECTURÍA SÁNCHEZ",
    "COLECTURÍA SANTIAGO RODRÍGUEZ", "COLECTURÍA SOSUA PUERTO PLATA", "COLECTURÍA VILLA VÁSQUEZ",
    "DEPTO. COOPERACION INTERNACIONAL", "DEPTO. EDUCACION TRIBUTARIA.",
    "DEPTO. LIBRE ACCESO A LA INFORMACION PUBLICA", "DEPTO. RECAUDACION",
    "DEPTO. RECONSIDERACION", "DEPTO. SEGURIDAD ELECTRONICA", "DEPTO. VALORACION DE BIENES",
    "ESTAFETA AEROPUERTO LAS AMÉRICAS ZONA ORIENTAL", "ESTAFETA LA FERIA",
    "ESTAFETA PUERTO CAUCEDO BOCA CHICA", "ESTAFETA PUERTO DE HAINA SAN CRISTÓBAL",
    "ESTAFETA PUERTO PLATA", "ESTAFETA PUERTO SANSOUCI ZONA ORIENTAL",
    "ESTAFETA PUERTO SANTO DOMINGO SAN CARLOS",
    "GERENCIA ADMINISTRATIVA Y FINANCIERA – Depto. Administración de Inventarios",
    "GERENCIA ADMINISTRATIVA Y FINANCIERA – Depto. Compras",
    "GERENCIA  ADMINISTRATIVA Y FINANCIERA – Depto. Contabilidad",
    "GERENCIA  ADMINISTRATIVA Y FINANCIERA – Depto. Tesorería",
    "GCIA. APOYO LOGISTICO - Depto. De Ingeniería", "GCIA. APOYO LOGISTICO - Transportación",
    "GCIA. COBRANZA – COBRO COACTIVO", "GCIA. COBRANZA – COBRO PERSUASIVO",
    "GCIA. FISCALIZACION EXTERNA GRANDES CONTRIBUYENTES",
    "GCIA. FISCALIZACION EXTERNA MEDIANOS CONTRIBUYENTES",
    "GCIA. II FISCALIZACION EXTERNA DE MEDIANOS CONTRIBUYENTES",
    "GCIA. INVESTIGACION DE FRAUDES Y DELITOS TRIBUTARIOS",
    "GCIA. PLANIFICACION CONTROL TRIBUTARIO - Depto. Planes selectivos",
    "GCIA. PLANIFICACION CONTROL TRIBUTARIO - Planes masivos",
    "GCIA. PLANIFICACION ESTRATEGICA", "GERENCIA COMUNICACIÓN ESTRATEGICA",
    "GERENCIA CONTROL RIESGOS - Depto. Auditoría Interna",
    "GERENCIA CONTROL RIESGOS - Depto. Control Previo Financiero",
    "GERENCIA CONTROL RIESGOS - Depto. Investigación Denuncias", "GERENCIA DE FACTURACION",
    "GERENCIA DE GRANDES CONTRIBUYENTES",
    "GERENCIA DE GRANDES CONTRIBUYENTES - Depto. Crédito y Compensaciones",
    "GERENCIA DE GRANDES CONTRIBUYENTES - Depto. Fiscalización Interna",
    "GERENCIA DE GRANDES CONTRIBUYENTES - Depto. Impuesto Selectivo al Consumo",
    "GERENCIA DE GRANDES CONTRIBUYENTES - Depto. Precio de Transferencia"
]

def normalizar_area(area_raw):
    if not area_raw: return None
    area_raw = str(area_raw).strip().upper()
    matches = get_close_matches(area_raw, [a.upper() for a in CATALOGO_AREAS], n=1, cutoff=0.5)
    return CATALOGO_AREAS[[a.upper() for a in CATALOGO_AREAS].index(matches[0])] if matches else area_raw

# ==========================================
# FUNCIONES DE LECTURA ABSOLUTA
# ==========================================
def celda_val(ws, col, row):
    """Lectura estricta por coordenadas. Nada de escáneres que corren columnas."""
    try:
        val = ws[f"{col}{row}"].value
        if val is None: return None
        if isinstance(val, datetime.datetime): return val.strftime("%d/%m/%Y")
        val_str = str(val).strip()
        if not val_str or val_str in ['-', '.', '_', '=', ' ']: return None
        return val_str
    except: return None

def limpiar_fecha_excel(fecha_raw):
    if not fecha_raw: return None
    fecha_str = str(fecha_raw).strip()
    if fecha_str.isdigit() and 30000 < int(fecha_str) < 80000:
        try: return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=int(fecha_str))).strftime("%d/%m/%Y")
        except: pass
    return fecha_str

# ==========================================
# 🟢 CAPA 1: EXPLORADOR
# ==========================================
class CapaExplorador:
    @staticmethod
    def ejecutar():
        print("\n=== INICIANDO CAPA 1: EXPLORACIÓN DE ARCHIVOS ===")
        root = tk.Tk()
        root.withdraw()
        
        opcion = input("1. Carpeta con múltiples archivos\n2. Un solo archivo\nSeleccione: ").strip()
        archivos = []
        
        if opcion == '1':
            carpeta = filedialog.askdirectory(title="Seleccione la carpeta")
            if carpeta:
                print(f"Buscando en: {carpeta}")
                for root_dir, _, files in os.walk(carpeta):
                    for file in files:
                        if file.endswith('.xlsx') and not file.startswith('~$'):
                            archivos.append(os.path.join(root_dir, file))
        else:
            archivo = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])
            if archivo: archivos.append(archivo)
                
        print(f"-> Archivos Excel encontrados: {len(archivos)}")
        return archivos

# ==========================================
# 🟡 CAPA 2: VALIDADOR POR COORDENADAS
# ==========================================
class CapaValidador:
    @staticmethod
    def identificar_portada(ws):
        # Según tu documento, revisamos celdas fijas D12 y D9
        d12 = str(celda_val(ws, "D", 12)).upper()
        d9 = str(celda_val(ws, "D", 9)).upper()
        
        if "FORMULARIO PARA TRANSFERENCIA" in d12: return "PORTADA_RECIENTE"
        if "FORMULARIO PARA TRANSFERENCIA" in d9: return "PORTADA_ANTIGUA"
        return None

    @staticmethod
    def identificar_anexo(ws, version_portada):
        titulo = ws.title.upper()
        if "1 CAJA" in titulo or "UNA CAJA" in titulo: return "1_CAJA"
        if "VARIAS CAJAS" in titulo: return "VARIAS_CAJAS"
        
        # Respaldo basado en tus coordenadas exactas
        if version_portada == "PORTADA_RECIENTE":
            if "CAJA" in str(celda_val(ws, "F", 27)).upper(): return "VARIAS_CAJAS"
            if "CÉDULA" in str(celda_val(ws, "F", 28)).upper() or "CEDULA" in str(celda_val(ws, "F", 28)).upper(): return "1_CAJA"
        elif version_portada == "PORTADA_ANTIGUA":
            if "CAJA" in str(celda_val(ws, "F", 25)).upper(): return "VARIAS_CAJAS"
            if "CÉDULA" in str(celda_val(ws, "F", 27)).upper() or "CEDULA" in str(celda_val(ws, "F", 27)).upper(): return "1_CAJA"
            
        return "IGNORAR"

    @staticmethod
    def ejecutar(archivos):
        print("\n=== INICIANDO CAPA 2: LECTURA DE METADATOS ===")
        plan_ejecucion = {} 
        
        for ruta in archivos:
            try:
                wb = load_workbook(ruta, data_only=True)
                version = None
                portada_data = None
                
                # Buscar y extraer Portada
                for ws in wb.worksheets:
                    version = CapaValidador.identificar_portada(ws)
                    if version == "PORTADA_RECIENTE":
                        dia, mes, anio = celda_val(ws, "L", 15), celda_val(ws, "N", 15), celda_val(ws, "O", 15)
                        portada_data = {
                            'version': version,
                            'fecha_recepcion': f"{dia}/{mes}/{anio}" if dia and mes and anio else None,
                            'area': normalizar_area(celda_val(ws, "G", 22)),
                            'motivo': celda_val(ws, "G", 27),
                            'cajas_mapeo': {}
                        }
                        for r in [31, 33, 35, 37, 39, 41, 43, 45]:
                            doc, caja = celda_val(ws, "G", r), celda_val(ws, "O", r)
                            if doc and caja: portada_data['cajas_mapeo'][str(caja).strip().upper()] = doc
                        break
                        
                    elif version == "PORTADA_ANTIGUA":
                        dia, mes, anio = celda_val(ws, "L", 12), celda_val(ws, "N", 12), celda_val(ws, "O", 12)
                        portada_data = {
                            'version': version,
                            'fecha_recepcion': f"{dia}/{mes}/{anio}" if dia and mes and anio else None,
                            'area': normalizar_area(celda_val(ws, "G", 19)),
                            'motivo': None,
                            'cajas_mapeo': {}
                        }
                        for r in [28, 30, 32, 34, 36, 38, 40, 42]:
                            doc, caja = celda_val(ws, "G", r), celda_val(ws, "O", r)
                            if doc and caja: portada_data['cajas_mapeo'][str(caja).strip().upper()] = doc
                        break
                
                if not portada_data:
                    wb.close()
                    continue 
                    
                # Evaluar Hojas Anexas
                hojas_validas = []
                for ws in wb.worksheets:
                    tipo_anexo = CapaValidador.identificar_anexo(ws, version)
                    if tipo_anexo in ["1_CAJA", "VARIAS_CAJAS"]:
                        hojas_validas.append({'nombre_hoja': ws.title, 'tipo': tipo_anexo})
                        
                if hojas_validas:
                    plan_ejecucion[ruta] = {'metadatos': portada_data, 'hojas': hojas_validas}
                    
                wb.close()
            except Exception: pass
                
        print(f"-> Archivos válidos para procesar: {len(plan_ejecucion)}")
        return plan_ejecucion

# ==========================================
# 🔵 CAPA 3: MIGRACIÓN ESTRICTA
# ==========================================
class CapaMigrador:
    @staticmethod
    def coincidir_caja(caja_encontrada, catalogo_cajas):
        if not caja_encontrada: return list(catalogo_cajas.values())[0] if len(catalogo_cajas) == 1 else None
        c = str(caja_encontrada).strip().upper()
        if c in catalogo_cajas: return catalogo_cajas[c]
        for key, doc in catalogo_cajas.items():
            if key in c or c in key: return doc
        return list(catalogo_cajas.values())[0] if len(catalogo_cajas) == 1 else None

    @staticmethod
    def ejecutar(plan_ejecucion, db_path):
        print("\n=== INICIANDO CAPA 3: EXTRACCIÓN MILIMÉTRICA ===")
        registros_a_migrar = []
        errores_reales = []

        for ruta, plan in plan_ejecucion.items():
            wb = load_workbook(ruta, data_only=True)
            meta = plan['metadatos']
            version = meta['version']
            cajas_mapeo = meta['cajas_mapeo']

            for hoja in plan['hojas']:
                ws = wb[hoja['nombre_hoja']]
                tipo = hoja['tipo']
                
                if tipo == "1_CAJA":
                    if version == "PORTADA_RECIENTE":
                        caja_global = celda_val(ws, "F", 20)
                        inicio, fin = 29, 128
                    else:
                        caja_global = celda_val(ws, "F", 19)
                        inicio, fin = 28, 127
                        
                    tipo_doc = CapaMigrador.coincidir_caja(caja_global, cajas_mapeo)
                    
                    for r in range(inicio, fin + 1):
                        cedula = celda_val(ws, "F", r)
                        nombre = celda_val(ws, "G", r)
                        fecha = limpiar_fecha_excel(celda_val(ws, "J", r))
                        
                        if not cedula and not nombre: continue 
                        
                        if not all([cedula, nombre, fecha, caja_global, tipo_doc]):
                            errores_reales.append([ruta, ws.title, r, "DATO INCOMPLETO", "Falta Cédula, Nombre, Fecha o Mapeo de Caja"])
                            continue
                            
                        registros_a_migrar.append({
                            "FECHA_RECEPCION": meta['fecha_recepcion'], "AREA": meta['area'], "CAJA": caja_global,
                            "CEDULA": cedula, "NOMBRE": nombre, "TIPO_DOC": tipo_doc, "FECHA": fecha, "MOTIVO": meta['motivo']
                        })

                elif tipo == "VARIAS_CAJAS":
                    if version == "PORTADA_RECIENTE":
                        inicio, fin = 28, 127
                    else:
                        inicio, fin = 26, 125
                        
                    for r in range(inicio, fin + 1):
                        caja = celda_val(ws, "F", r)
                        cedula = celda_val(ws, "G", r)
                        nombre = celda_val(ws, "H", r)
                        fecha = limpiar_fecha_excel(celda_val(ws, "K", r))
                        
                        if not cedula and not nombre: continue 
                        
                        tipo_doc = CapaMigrador.coincidir_caja(caja, cajas_mapeo)
                        
                        if not all([caja, cedula, nombre, fecha, tipo_doc]):
                            errores_reales.append([ruta, ws.title, r, "DATO INCOMPLETO", "Falta Caja, Cédula, Nombre, Fecha o Mapeo de Caja"])
                            continue
                            
                        registros_a_migrar.append({
                            "FECHA_RECEPCION": meta['fecha_recepcion'], "AREA": meta['area'], "CAJA": caja,
                            "CEDULA": cedula, "NOMBRE": nombre, "TIPO_DOC": tipo_doc, "FECHA": fecha, "MOTIVO": meta['motivo']
                        })
            wb.close()

        print("💾 Escribiendo Base de Datos...")
        if registros_a_migrar:
            db = load_workbook(db_path)
            ws_bd = db["Base de Datos"]
            tabla = ws_bd.tables["BD"]
            start_col, start_row, end_col, end_row = range_boundaries(tabla.ref)
            
            insert_row = start_row + 1
            while ws_bd.cell(row=insert_row, column=start_col).value is not None: insert_row += 1
                
            for rec in registros_a_migrar:
                ws_bd.cell(row=insert_row, column=1).value = rec.get("FECHA_RECEPCION")
                ws_bd.cell(row=insert_row, column=2).value = rec.get("AREA")
                ws_bd.cell(row=insert_row, column=3).value = rec.get("CAJA")
                ws_bd.cell(row=insert_row, column=4).value = "N/A"
                ws_bd.cell(row=insert_row, column=5).value = "N/A"
                ws_bd.cell(row=insert_row, column=6).value = "N/A"
                ws_bd.cell(row=insert_row, column=7).value = ""
                ws_bd.cell(row=insert_row, column=8).value = rec.get("CEDULA")
                ws_bd.cell(row=insert_row, column=9).value = rec.get("NOMBRE")
                ws_bd.cell(row=insert_row, column=10).value = rec.get("TIPO_DOC")
                ws_bd.cell(row=insert_row, column=11).value = rec.get("FECHA")
                ws_bd.cell(row=insert_row, column=12).value = rec.get("MOTIVO")
                # Las columnas M, N, O, P quedan nulas (en blanco) por defecto
                insert_row += 1
                
            tabla.ref = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{insert_row - 1}"
            db.save(db_path)
            db.close()

        if errores_reales:
            wb_err = Workbook()
            ws_err = wb_err.active
            ws_err.append(["RUTA_ARCHIVO", "HOJA", "FILA", "TIPO_ERROR", "DESCRIPCION"])
            for err in errores_reales: ws_err.append(err)
            wb_err.save("Errores_Migracion_Formularios.xlsx")
            wb_err.close()

        print(f"\n✅ Proceso Finalizado")
        print(f"Registros migrados: {len(registros_a_migrar)}")
        print(f"Datos incompletos no migrados: {len(errores_reales)}")

# ==========================================
# 🚀 ORQUESTADOR
# ==========================================
if __name__ == "__main__":
    db_path = "Base de Datos Control de Documentos.xlsx"
    if not os.path.exists(db_path):
        print(f"❌ ERROR: Falta archivo '{db_path}'.")
    else:
        archivos = CapaExplorador.ejecutar()
        if archivos:
            plan = CapaValidador.ejecutar(archivos)
            if plan: CapaMigrador.ejecutar(plan, db_path)